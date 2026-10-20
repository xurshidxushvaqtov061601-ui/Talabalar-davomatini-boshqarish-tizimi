import io
from fastapi import APIRouter, Depends, Response
from sqlalchemy.orm import Session
from sqlalchemy import func
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.database import get_db
from app.models import Group, Faculty, Student, Attendance, Lesson, Subject
from app.auth import get_current_user, User

router = APIRouter(prefix="/api/reports", tags=["Reports"])

@router.get("/summary")
def get_reports_summary(db: Session = Depends(get_db)):
    groups = db.query(Group).all()
    group_stats = []

    for g in groups:
        students = g.students
        student_ids = [s.id for s in students]
        total_records = db.query(Attendance).filter(Attendance.student_id.in_(student_ids)).count() if student_ids else 0
        present = db.query(Attendance).filter(Attendance.student_id.in_(student_ids), Attendance.status == "keldi").count() if student_ids else 0
        late = db.query(Attendance).filter(Attendance.student_id.in_(student_ids), Attendance.status == "kechikdi").count() if student_ids else 0
        absent = db.query(Attendance).filter(Attendance.student_id.in_(student_ids), Attendance.status.in_(["sababsiz", "sababli"])).count() if student_ids else 0

        rate = round(((present + late * 0.8) / total_records * 100), 1) if total_records > 0 else 100.0
        group_stats.append({
            "group_id": g.id,
            "group_name": g.name,
            "faculty_name": g.faculty.name if g.faculty else "",
            "students_count": len(students),
            "total_records": total_records,
            "present": present,
            "late": late,
            "absent": absent,
            "rate": rate
        })

    return group_stats

@router.get("/at-risk-students")
def get_at_risk_students(db: Session = Depends(get_db)):
    students = db.query(Student).all()
    at_risk = []

    for s in students:
        total = db.query(Attendance).filter(Attendance.student_id == s.id).count()
        if total >= 3:
            absent = db.query(Attendance).filter(
                Attendance.student_id == s.id,
                Attendance.status.in_(["sababsiz", "sababli"])
            ).count()
            unexcused = db.query(Attendance).filter(
                Attendance.student_id == s.id,
                Attendance.status == "sababsiz"
            ).count()
            absent_pct = round((absent / total * 100), 1)

            if absent_pct >= 20.0:
                at_risk.append({
                    "student_id": s.id,
                    "full_name": s.user.full_name if s.user else "",
                    "student_id_number": s.student_id_number,
                    "group_name": s.group.name if s.group else "",
                    "faculty_name": s.group.faculty.name if s.group and s.group.faculty else "",
                    "phone": s.user.phone if s.user else "",
                    "total_lessons": total,
                    "absent_lessons": absent,
                    "unexcused_lessons": unexcused,
                    "absent_percentage": absent_pct
                })

    at_risk.sort(key=lambda x: x["absent_percentage"], reverse=True)
    return at_risk

@router.get("/export-excel")
def export_excel(group_id: int = None, db: Session = Depends(get_db)):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Davomat Hisoboti"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    border_thin = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    headers = ["№", "FIO", "Talaba ID", "Guruh", "Fakultet", "Jami darslar", "Keldi", "Kechikdi", "Sababli", "Sababsiz", "Davomat %"]
    ws.append(headers)

    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    query = db.query(Student)
    if group_id:
        query = query.filter(Student.group_id == group_id)
    students = query.all()
    students.sort(key=lambda s: s.user.full_name if s.user else "")

    for idx, s in enumerate(students, 1):
        atts = db.query(Attendance).filter(Attendance.student_id == s.id).all()
        total = len(atts)
        present = sum(1 for a in atts if a.status == "keldi")
        late = sum(1 for a in atts if a.status == "kechikdi")
        excused = sum(1 for a in atts if a.status == "sababli")
        unexcused = sum(1 for a in atts if a.status == "sababsiz")

        pct = round(((present + late * 0.8) / total * 100), 1) if total > 0 else 100.0

        row = [
            idx,
            s.user.full_name if s.user else "",
            s.student_id_number,
            s.group.name if s.group else "",
            s.group.faculty.name if s.group and s.group.faculty else "",
            total,
            present,
            late,
            excused,
            unexcused,
            f"{pct}%"
        ]
        ws.append(row)

        row_num = idx + 1
        for col_num in range(1, len(headers) + 1):
            c = ws.cell(row=row_num, column=col_num)
            c.border = border_thin
            if col_num in [1, 3, 6, 7, 8, 9, 10, 11]:
                c.alignment = center_align

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=davomat_hisoboti.xlsx"}
    )
